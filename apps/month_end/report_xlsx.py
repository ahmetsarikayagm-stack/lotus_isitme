"""Ay Sonu Excel raporları: logolu başlık bandı, formüllü toplamlar (Arial), alt bilgide firma bilgileri.

Yıllık ve aylık rapor AYNI düzendedir: Kira, Elektrik, Su, İnternet, Maaş, Muhasebe, Genel Gid.,
Cihaz Gid., Top. Gid., Top. Gel., Kâr. Aylık dosyada ek olarak o ayın 'Satışlar' ve 'Kalemler' sayfaları bulunur.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.properties import CalcProperties
from PIL import Image as PILImage

from .reports import LOGO_WHITE, company_info

BRAND, BRAND_DARK, BRAND_LIGHT, GREY = "EF7B18", "B4510C", "FEF4EA", "94A3B8"
INPUT_BLUE = "0000FF"   # Ay Sonu kaydından alınan giriş değerleri

MONEY = '#,##0.00 "₺";[Red]-#,##0.00 "₺";"-"'
INTEGER = '#,##0;-#,##0;"-"'
DATE = "dd.mm.yyyy"

THIN = Side(style="thin", color="E2E8F0")
MEDIUM_BRAND = Side(style="medium", color=BRAND)

HEADERS = ["Ay", "Kira", "Elektrik", "Su", "İnternet", "Maaş", "Muhasebe", "Genel Gid.", "Cihaz Gid.", "Top. Gid.", "Top. Gel.", "Kâr"]


def _w(ws, ref, value=None, *, fmt=None, bold=False, color="000000", fill=None, size=10, h=None,
       wrap=False, border=None, italic=False, v="center"):
    cell = ws[ref]
    cell.value = value
    cell.font = Font(name="Arial", size=size, bold=bold, color=color, italic=italic)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if border:
        cell.border = border
    return cell


def _band(ws, title, last_col, prepared_by):
    """Üç satırlık turuncu başlık bandı: solda logo, yanında firma ve rapor bilgisi."""
    info = company_info()
    for row in (1, 2, 3):
        ws.row_dimensions[row].height = 24
        for col in range(1, last_col + 1):
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BRAND)
    ws.column_dimensions["A"].width = 15

    pil = PILImage.open(LOGO_WHITE)
    img = XLImage(str(LOGO_WHITE))
    height = 80
    width = int(height * pil.width / pil.height)
    img.height, img.width = height, width
    img.anchor = OneCellAnchor(                      # bandın kenarına yapışmasın: küçük iç boşluk
        _from=AnchorMarker(col=0, row=0, colOff=pixels_to_EMU(9), rowOff=pixels_to_EMU(5)),
        ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)),
    )
    ws.add_image(img)

    _w(ws, "B1", info["name"], bold=True, color="FFFFFF", size=16, fill=BRAND)
    _w(ws, "B2", title, bold=True, color="FFFFFF", size=12, fill=BRAND)
    _w(ws, "B3", f"Rapor tarihi: {datetime.now():%d.%m.%Y %H:%M}  |  Hazırlayan: {prepared_by}", color="FFFFFF", size=9, fill=BRAND)
    ws.sheet_view.showGridLines = False


def _widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _page(ws, landscape=True):
    """Yazdırma ayarları + yazdırma alt bilgisi (adres, telefon, e-posta, sayfa numarası)."""
    info = company_info()
    contact = "  ·  ".join(x for x in (info["phone"], info["email"]) if x)
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    # Not: yazdırma alt bilgisinde satır sonu karakteri güvenilir değil (_x000a_ olarak görünür); tek satır kullanılır
    ws.oddFooter.left.text = f"{info['address']}  ·  {contact}".replace("&", "&&")
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Sayfa &P / &N"
    ws.oddFooter.right.size = 8


def _footer_rows(ws, row, last_col):
    """Sayfa içinde de görünen alt bilgi satırları: adres ve telefon · e-posta."""
    info = company_info()
    contact = "  ·  ".join(x for x in (info["phone"], info["email"]) if x)
    for offset, text in enumerate((info["address"], contact)):
        ref = f"B{row + offset}"
        _w(ws, ref, text, color="64748B", size=9, h="center")
        ws.merge_cells(f"B{row + offset}:{get_column_letter(last_col)}{row + offset}")
    for col in range(2, last_col + 1):
        ws.cell(row=row, column=col).border = Border(top=THIN)


def _header_row(ws, row, labels, start_col=2, right_from=None):
    for i, text in enumerate(labels):
        col = get_column_letter(start_col + i)
        right = right_from is not None and i >= right_from
        _w(ws, f"{col}{row}", text, bold=True, color="FFFFFF", fill=BRAND_DARK, h="right" if right else "left", wrap=True)
    ws.row_dimensions[row].height = 30


def _workbook(title):
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)   # Excel açıldığında formülleri hesaplar
    wb.properties.creator = company_info()["name"]
    wb.properties.title = title
    return wb


def _bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================= ORTAK ANA TABLO (yıllık = aylık düzeni)
def _report_sheet(ws, title, rows, with_total, prepared_by, status=None):
    """
    Sütunlar: B Ay | C Kira | D Elektrik | E Su | F İnternet | G Maaş | H Muhasebe | I Genel Gid. |
              J Cihaz Gid. | K Top. Gid. (=SUM C:J) | L Top. Gel. | M Kâr (=L-K)
    Mavi değerler Ay Sonu kaydından gelir; K, M ve toplamlar formüldür.
    """
    _band(ws, title, 14, prepared_by)
    _widths(ws, {"B": 16, "C": 14, "D": 14, "E": 12, "F": 14, "G": 14, "H": 14, "I": 14, "J": 15, "K": 16, "L": 16, "M": 16, "N": 12})
    if status:
        _w(ws, "B4", status, color="0F172A" if "KAPALI" in status else "64748B", bold="KAPALI" in status, italic="KAPALI" not in status, size=9)
    _header_row(ws, 5, HEADERS, right_from=1)

    first = 6
    for i, row in enumerate(rows):
        r = first + i
        m = row["obj"]
        _w(ws, f"B{r}", row["name"], bold=True, border=Border(bottom=THIN))
        if m is None:
            for col in "CDEFGHIJKLM":
                _w(ws, f"{col}{r}", None, border=Border(bottom=THIN))
            _w(ws, f"N{r}", "Kayıt yok", color=GREY, italic=True, border=Border(bottom=THIN))
            continue
        inputs = [m.rent, m.electricity, m.water, m.internet, m.salary, m.accounting, m.general_expenses, m.device_cost]
        for col, value in zip("CDEFGHIJ", inputs):
            _w(ws, f"{col}{r}", float(value), fmt=MONEY, color=INPUT_BLUE, border=Border(bottom=THIN))
        _w(ws, f"K{r}", f"=SUM(C{r}:J{r})", fmt=MONEY, bold=True, border=Border(bottom=THIN))
        _w(ws, f"L{r}", float(m.total_income), fmt=MONEY, bold=True, color=INPUT_BLUE, border=Border(bottom=THIN))
        _w(ws, f"M{r}", f"=L{r}-K{r}", fmt=MONEY, bold=True, border=Border(bottom=THIN))
        _w(ws, f"N{r}", None, border=Border(bottom=THIN))

    last = first + len(rows) - 1
    nxt = last + 1
    if with_total:
        t = last + 1
        _w(ws, f"B{t}", "TOPLAM", bold=True, fill=BRAND_LIGHT)
        for col in "CDEFGHIJKLM":
            _w(ws, f"{col}{t}", f"=SUM({col}{first}:{col}{last})", fmt=MONEY, bold=True, fill=BRAND_LIGHT)
        _w(ws, f"N{t}", None, fill=BRAND_LIGHT)
        for col in "BCDEFGHIJKLMN":
            ws[f"{col}{t}"].border = Border(top=Side(style="medium", color="0F172A"))
        nxt = t + 1

    note = ("Mavi renkli değerler Ay Sonu kaydından alınmıştır; Top. Gid. (Kira + Elektrik + Su + İnternet + Maaş + Muhasebe + "
            "Genel Gid. + Cihaz Gid.), Kâr (Top. Gel. − Top. Gid.) ve toplamlar formülle hesaplanır."
            + (" Ay sonu kaydı girilmemiş aylar toplamlara dahil değildir." if with_total else ""))
    _w(ws, f"B{nxt + 1}", note, color="64748B", italic=True, size=9, wrap=True, v="top")
    ws.merge_cells(f"B{nxt + 1}:N{nxt + 1}")
    ws.row_dimensions[nxt + 1].height = 30
    _footer_rows(ws, nxt + 3, 14)
    ws.freeze_panes = "C6"
    _page(ws)


# ================================================================= YILLIK RAPOR
def build_year_xlsx(data: dict, prepared_by: str = "") -> bytes:
    year = data["year"]
    wb = _workbook(f"Yıllık Ay Sonu Raporu – {year}")
    ws = wb.active
    ws.title = "Yıllık Rapor"
    _report_sheet(ws, f"Yıllık Ay Sonu Raporu – {year}", [{"name": r["name"], "obj": r["obj"]} for r in data["rows"]],
                  with_total=True, prepared_by=prepared_by)
    return _bytes(wb)


# ================================================================= AYLIK RAPOR (yıllıkla aynı düzen + destek sayfaları)
def _status_line(me):
    if me.is_closed:
        who = f" · {me.closed_by.display_name}" if me.closed_by else ""
        when = f"{me.closed_at:%d.%m.%Y %H:%M}" if me.closed_at else ""
        return f"KAPALI DÖNEM — {when}{who} tarafından kapatılmıştır; değerler dondurulmuştur."
    return "Dönem açıktır; değerler güncel satışlara göre hesaplanmıştır."


def build_month_xlsx(data: dict, prepared_by: str = "") -> bytes:
    me = data["me"]
    sales, items = data["sales"], data["items"]

    wb = _workbook(f"Ay Sonu Raporu – {me.period_label}")
    ws = wb.active
    ws.title = "Ay Sonu"
    ws_sales = wb.create_sheet("Satışlar")
    ws_items = wb.create_sheet("Kalemler")

    _report_sheet(ws, f"Ay Sonu Raporu – {me.period_label}", [{"name": me.get_month_display(), "obj": me}],
                  with_total=False, prepared_by=prepared_by, status=_status_line(me))

    # ---------------------------------------------------------- Satışlar
    _band(ws_sales, f"Satış Listesi – {me.period_label}", 13, prepared_by)
    _widths(ws_sales, {"B": 12, "C": 26, "D": 14, "E": 44, "F": 16, "G": 14, "H": 16, "I": 16, "J": 16, "K": 16, "L": 18, "M": 8})
    _header_row(ws_sales, 5, ["Tarih", "Hasta", "T.C. No", "Ürünler", "Ara Toplam", "İndirim", "Toplam",
                              "Kapora", "Kalan", "Ödeme", "Banka", "SGK"])
    for col in "FGHIJ":
        ws_sales[f"{col}5"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    for col in "KL":
        ws_sales[f"{col}5"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s_first = 6
    for i, sale in enumerate(sales):
        r = s_first + i
        parts = [f"{it.product}" + (f" ×{it.quantity}" if it.quantity > 1 else "") for it in sale.items.all()]
        _w(ws_sales, f"B{r}", sale.sale_date, fmt=DATE, h="left")
        _w(ws_sales, f"C{r}", str(sale.patient))
        _w(ws_sales, f"D{r}", sale.patient.tc_no, h="left")
        _w(ws_sales, f"E{r}", " + ".join(parts), wrap=True)
        _w(ws_sales, f"F{r}", float(sale.subtotal), fmt=MONEY)
        _w(ws_sales, f"G{r}", float(sale.discount), fmt=MONEY)
        _w(ws_sales, f"H{r}", f"=F{r}-G{r}", fmt=MONEY, bold=True)
        _w(ws_sales, f"I{r}", float(sale.deposit), fmt=MONEY)
        _w(ws_sales, f"J{r}", f"=H{r}-I{r}", fmt=MONEY)
        _w(ws_sales, f"K{r}", sale.payment_method.name)
        _w(ws_sales, f"L{r}", sale.bank.name if sale.bank else "")
        _w(ws_sales, f"M{r}", "Evet" if sale.has_institution else "Hayır", h="center")
        for col in "KL":
            ws_sales[f"{col}{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in "BCDEFGHIJKLM":
            ws_sales[f"{col}{r}"].border = Border(bottom=THIN)
    s_last = s_first + max(len(sales), 1) - 1
    t = s_last + 1
    _w(ws_sales, f"B{t}", "TOPLAM", bold=True, fill=BRAND_LIGHT)
    for col in "CDEKLM":
        _w(ws_sales, f"{col}{t}", None, fill=BRAND_LIGHT)
    for col in "FGHIJ":
        _w(ws_sales, f"{col}{t}", f"=SUM({col}{s_first}:{col}{s_last})", fmt=MONEY, bold=True, fill=BRAND_LIGHT)
    for col in "BCDEFGHIJKLM":
        ws_sales[f"{col}{t}"].border = Border(top=Side(style="medium", color="0F172A"))
    ws_sales.auto_filter.ref = f"B5:M{s_last}"
    ws_sales.freeze_panes = "A6"
    _footer_rows(ws_sales, t + 2, 13)
    _page(ws_sales)

    # ---------------------------------------------------------- Kalemler
    _band(ws_items, f"Satış Kalemleri – {me.period_label}", 13, prepared_by)
    _widths(ws_items, {"B": 12, "C": 26, "D": 16, "E": 16, "F": 24, "G": 8, "H": 16, "I": 17, "J": 16, "K": 17, "L": 16, "M": 8})
    _header_row(ws_items, 5, ["Tarih", "Hasta", "Kategori", "Marka", "Model", "Adet", "Birim Fiyat", "Satır Toplamı",
                              "Birim Maliyet", "Maliyet Toplamı", "Brüt Kâr", "SGK"])
    for col in "GHIJKL":
        ws_items[f"{col}5"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    i_first = 6
    for i, it in enumerate(items):
        r = i_first + i
        _w(ws_items, f"B{r}", it.sale.sale_date, fmt=DATE, h="left")
        _w(ws_items, f"C{r}", str(it.sale.patient))
        _w(ws_items, f"D{r}", it.category.name)
        _w(ws_items, f"E{r}", it.product.brand.name)
        _w(ws_items, f"F{r}", it.product.model.name)
        _w(ws_items, f"G{r}", it.quantity, fmt=INTEGER)
        _w(ws_items, f"H{r}", float(it.unit_price), fmt=MONEY)
        _w(ws_items, f"I{r}", f"=G{r}*H{r}", fmt=MONEY, bold=True)
        _w(ws_items, f"J{r}", float(it.unit_cost), fmt=MONEY)
        _w(ws_items, f"K{r}", f"=G{r}*J{r}", fmt=MONEY)
        _w(ws_items, f"L{r}", f"=I{r}-K{r}", fmt=MONEY)
        _w(ws_items, f"M{r}", "Evet" if it.institution_right else "Hayır", h="center")
        for col in "BCDEFGHIJKLM":
            ws_items[f"{col}{r}"].border = Border(bottom=THIN)
    i_last = i_first + max(len(items), 1) - 1
    t2 = i_last + 1
    _w(ws_items, f"B{t2}", "TOPLAM", bold=True, fill=BRAND_LIGHT)
    for col in "CDEFHJM":
        _w(ws_items, f"{col}{t2}", None, fill=BRAND_LIGHT)
    for col, fmt in (("G", INTEGER), ("I", MONEY), ("K", MONEY), ("L", MONEY)):
        _w(ws_items, f"{col}{t2}", f"=SUM({col}{i_first}:{col}{i_last})", fmt=fmt, bold=True, fill=BRAND_LIGHT)
    for col in "BCDEFGHIJKLM":
        ws_items[f"{col}{t2}"].border = Border(top=Side(style="medium", color="0F172A"))
    ws_items.auto_filter.ref = f"B5:M{i_last}"
    ws_items.freeze_panes = "A6"
    _footer_rows(ws_items, t2 + 2, 13)
    _page(ws_items)
    return _bytes(wb)
